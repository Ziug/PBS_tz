import os, sys, json
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from fastapi import FastAPI, Request, Form, HTTPException
from fastapi.responses import HTMLResponse, RedirectResponse
from starlette.middleware.sessions import SessionMiddleware
from fastapi.templating import Jinja2Templates
from fastapi.staticfiles import StaticFiles
from datetime import datetime, date as dt_date

def parse_date_input(v):
    v = (v or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(v, fmt).date()
        except ValueError:
            continue
    raise ValueError("bad date")
from sqlalchemy.orm import Session
from app.database import SessionLocal, init_db
from app.models import PrepRun, WeldNeed, DocRun, DataIssue
from app.ingest.quality import classify_all
from collections import defaultdict

app = FastAPI(title="PBS demo")

@app.middleware("http")
async def auth_middleware(request, call_next):
    path = request.url.path
    if path.startswith("/static") or path == "/login":
        return await call_next(request)
    if not request.session.get("user"):
        from fastapi.responses import RedirectResponse
        return RedirectResponse("/login", status_code=303)
    return await call_next(request)

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
templates = Jinja2Templates(directory=os.path.join(BASE_DIR, "templates"))
templates.env.globals["getattr"] = getattr
app.mount("/static", StaticFiles(directory=os.path.join(BASE_DIR, "static")), name="static")

SECTIONS = {
  "prep": {"model": PrepRun, "title": "Заготовка", "fields": ["record_id", "date", "batch_id", "employee", "plan_qty", "fact_qty"]},
  "weld": {"model": WeldNeed, "title": "Потребность сварки", "fields": ["record_id", "date", "batch_id", "need_qty"]},
  "documents": {"model": DocRun, "title": "Документация", "fields": ["record_id", "date", "employee", "plan_docs", "fact_docs"]},
}

def db():
    return SessionLocal()


def next_record_id(model, prefix):
    try:
        from sqlalchemy import func
        import re
        s = db()
        rows = s.query(model).all()
        maxn = 0
        for r in rows:
            m = re.search(r"(\d+)$", str(r.record_id))
            if m:
                maxn = max(maxn, int(m.group(1)))
        s.close()
        return f"{prefix}-{str(maxn+1).zfill(3)}"
    except Exception:
        return f"{prefix}-001"

def current_user(request):
    return request.session.get("user")

def is_ceo(request):
    return request.session.get("role") == "ceo"

def require_login(request):
    if not current_user(request):
        raise HTTPException(status_code=303, headers={"Location": "/login"})

@app.on_event("startup")
def startup():
    init_db()


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    s = db()
    try:
        names = set()
        for r in s.query(PrepRun).all():
            if r.employee: names.add(r.employee)
        for r in s.query(DocRun).all():
            if r.employee: names.add(r.employee)
        names = sorted(names)
        return templates.TemplateResponse("login.html", {"request": request, "names": names})
    finally:
        s.close()

@app.post("/login")
async def login_do(request: Request):
    data = await request.form()
    name = data.get("user")
    if name == "__ceo__":
        request.session["user"] = "Владимир Мошкин"
        request.session["role"] = "ceo"
    elif name:
        request.session["user"] = name
        request.session["role"] = "employee"
    return RedirectResponse("/", status_code=303)

@app.get("/logout")
def logout(request: Request):
    request.session.clear()
    return RedirectResponse("/login", status_code=303)

@app.get("/", response_class=HTMLResponse)
def index(request: Request):
    return templates.TemplateResponse("index.html", {"request": request})

def get_issues_map(sec):
    table = {"prep": "prep_runs", "weld": "weld_needs", "documents": "docs_runs"}[sec]
    s = db()
    try:
        issues = s.query(DataIssue).filter_by(table_name=table).all()
        m = {}
        for i in issues:
            m.setdefault(i.record_id, []).append(i.message)
        return m
    finally:
        s.close()

@app.get("/{sec}", response_class=HTMLResponse)
def list_sec(request: Request, sec: str, q: str = "", sort: str = "id", dir: str = "asc"):
    if sec == "ceo":
        return dashboard(request, request.query_params.get("start",""), request.query_params.get("end",""))
    if sec == "export":
        return export_csv(request, request.query_params.get("start",""), request.query_params.get("end",""))
    if sec in ("export.xlsx","export_xlsx"):
        return export_xlsx(request, request.query_params.get("start",""), request.query_params.get("end",""))
    if sec not in SECTIONS:
        raise HTTPException(404)
    if sec == "weld" and not is_ceo(request):
        return templates.TemplateResponse("no_access.html", {"request": request}, status_code=403)
    cfg = SECTIONS[sec]
    M = cfg["model"]
    s = db()
    try:
        query = s.query(M)
        if (not is_ceo(request)) and sec in ("prep","documents"):
            query = query.filter(M.employee == current_user(request))
        if q:
            like = f"%{q}%"
            from sqlalchemy import or_
            cols = [getattr(M, c) for c in cfg["fields"] if c in ("record_id", "batch_id", "employee")]
            if cols:
                query = query.filter(or_(*[c.like(like) for c in cols]))
        rows = query.all()
        reverse = (dir == "desc")
        try:
            rows.sort(key=lambda r: getattr(r, sort), reverse=reverse)
        except Exception:
            pass
        issues = get_issues_map(sec)
        return templates.TemplateResponse("section.html", {"request": request, "sec": sec, "cfg": cfg,
                                                          "rows": rows, "issues": issues, "q": q, "sort": sort, "dir": dir})
    finally:
        s.close()

PREFIX = {"prep": "P", "weld": "W", "documents": "D"}
@app.get("/{sec}/new", response_class=HTMLResponse)
def new_form(request: Request, sec: str):
    if sec not in SECTIONS:
        raise HTTPException(404)
    cfg = SECTIONS[sec]
    M = cfg["model"]
    pr = PREFIX[sec]
    defaults = {"record_id": next_record_id(M, pr), "date": dt_date.today().isoformat()}
    if "employee" in cfg["fields"] and not is_ceo(request):
        # hide employee but still set server-side later
        pass
    return templates.TemplateResponse("form.html", {"request": request, "sec": sec, "cfg": cfg, "row": None,
                                                  "errors": [], "defaults": defaults, "is_ceo": is_ceo(request)})

def validate(sec, fields, data):
    errors = []
    for f in fields:
        if f in ("record_id", "date"):
            if not data.get(f):
                errors.append(f"Поле {f} обязательно")
        if f == "date":
            try:
                parse_date_input(data.get(f))
            except Exception:
                errors.append("Неверный формат даты (DD/MM/YYYY или YYYY-MM-DD)")
        if f in ("plan_qty", "fact_qty", "need_qty", "plan_docs", "fact_docs"):
            v = data.get(f)
            if v not in (None, ""):
                try:
                    iv = int(v)
                    if iv < 0:
                        errors.append(f"{f}: значение не может быть отрицательным")
                except Exception:
                    errors.append(f"{f}: должно быть числом")
    return errors

def apply_fields(row, fields, data, sec):
    for f in fields:
        v = data.get(f)
        if f == "date":
            setattr(row, f, parse_date_input(v))
        elif f in ("plan_qty", "fact_qty", "need_qty", "plan_docs", "fact_docs"):
            setattr(row, f, int(v) if v not in (None, "") else None)
        else:
            setattr(row, f, (v or "").strip())

@app.post("/{sec}/new")
async def create_async(request: Request, sec: str):
    if sec not in SECTIONS:
        raise HTTPException(404)
    cfg = SECTIONS[sec]
    M = cfg["model"]
    fields = cfg["fields"]
    data = await request.form()
    if (not is_ceo(request)) and "employee" in fields:
        # keep provided but override server-side
        pass
    errors = validate(sec, fields, data)
    if errors:
        return templates.TemplateResponse("form.html", {"request": request, "sec": sec, "cfg": cfg, "row": None, "errors": errors, "defaults": {}, "is_ceo": is_ceo(request), "data": data}, status_code=400)
    s = db()
    try:
        row = M()
        if (not is_ceo(request)) and "employee" in fields:
            data._mutable = True if hasattr(data, "_mutable") else None
        apply_fields(row, fields, data, sec)
        if (not is_ceo(request)) and "employee" in fields:
            row.employee = current_user(request)
        row.import_batch = "manual"
        row.source_row = 0
        if hasattr(row, "quality_status"):
            row.quality_status = "pending"
        s.add(row)
        s.flush()
        classify_all(s)
        s.commit()
        return RedirectResponse(f"/{sec}?msg=created", status_code=303)
    except Exception as e:
        s.rollback()
        return templates.TemplateResponse("form.html", {"request": request, "sec": sec, "cfg": cfg, "row": None, "errors": [str(e)]}, status_code=400)
    finally:
        s.close()

@app.get("/{sec}/{rid}/edit", response_class=HTMLResponse)
def edit_form(request: Request, sec: str, rid: int):
    if sec not in SECTIONS:
        raise HTTPException(404)
    cfg = SECTIONS[sec]
    M = cfg["model"]
    s = db()
    try:
        row = s.query(M).get(rid)
        if not row:
            raise HTTPException(404)
        if (not is_ceo(request)) and hasattr(row, "employee") and row.employee != current_user(request):
            raise HTTPException(403)
        return templates.TemplateResponse("form.html", {"request": request, "sec": sec, "cfg": cfg, "row": row, "errors": [], "defaults": {}, "is_ceo": is_ceo(request)})
    finally:
        s.close()

@app.post("/{sec}/{rid}/edit")
async def update(request: Request, sec: str, rid: int):
    if sec not in SECTIONS:
        raise HTTPException(404)
    cfg = SECTIONS[sec]
    M = cfg["model"]
    fields = cfg["fields"]
    data = await request.form()
    s = db()
    try:
        row = s.query(M).get(rid)
        if not row:
            raise HTTPException(404)
        if (not is_ceo(request)) and hasattr(row, "employee") and row.employee != current_user(request):
            raise HTTPException(403)
        errors = validate(sec, fields, data)
        if errors:
            return templates.TemplateResponse("form.html", {"request": request, "sec": sec, "cfg": cfg, "row": row, "errors": errors, "defaults": {}, "is_ceo": is_ceo(request), "data": data}, status_code=400)
        apply_fields(row, fields, data, sec)
        if (not is_ceo(request)) and "employee" in fields:
            row.employee = current_user(request)
        row.user_edited = True
        s.flush()
        classify_all(s)
        s.commit()
        return RedirectResponse(f"/{sec}?msg=saved", status_code=303)
    except Exception as e:
        s.rollback()
        return templates.TemplateResponse("form.html", {"request": request, "sec": sec, "cfg": cfg, "row": row, "errors": [str(e)], "defaults": {}, "is_ceo": is_ceo(request)}, status_code=400)
    finally:
        s.close()

@app.get("/ceo", response_class=HTMLResponse)
def dashboard(request: Request, start: str = "", end: str = ""):
    if not is_ceo(request):
        raise HTTPException(403, "Нет доступа")
    s = db()
    try:
        def parse(x):
            try:
                return dt_date.fromisoformat(x)
            except Exception:
                return None
        d0 = parse(start)
        d1 = parse(end)
        def filt(q, model):
            if d0:
                q = q.filter(model.date >= d0)
            if d1:
                q = q.filter(model.date <= d1)
            return q

        prep_rows = filt(s.query(PrepRun), PrepRun).all()
        weld_rows = filt(s.query(WeldNeed), WeldNeed).all()
        doc_rows = filt(s.query(DocRun), DocRun).all()

        prep_issues = s.query(DataIssue).filter_by(table_name="prep_runs", resolved=False).count()
        weld_issues = s.query(DataIssue).filter_by(table_name="weld_needs", resolved=False).count()
        doc_issues = s.query(DataIssue).filter_by(table_name="docs_runs", resolved=False).count()

        prep_plan = sum((r.plan_qty or 0) for r in prep_rows)
        prep_fact_rows = [r for r in prep_rows if r.fact_qty is not None]
        prep_fact = sum(r.fact_qty for r in prep_fact_rows)
        prep_total = len(prep_rows)
        prep_completion = (prep_fact / prep_plan) if prep_plan else None
        prep_partial = len(prep_fact_rows) < prep_total

        weld_need = sum((r.need_qty or 0) for r in weld_rows)
        weld_total = len(weld_rows)

        doc_plan = sum((r.plan_docs or 0) for r in doc_rows)
        doc_fact = sum((r.fact_docs or 0) for r in doc_rows)
        doc_total = len(doc_rows)
        doc_completion = (doc_fact / doc_plan) if doc_plan else None

        daily = defaultdict(lambda: [0, 0])
        for r in prep_rows:
            daily[str(r.date)][0] += r.plan_qty or 0
            if r.fact_qty is not None:
                daily[str(r.date)][1] += r.fact_qty
        chart_labels = sorted(daily.keys())
        chart_plan = [daily[d][0] for d in chart_labels]
        chart_fact = [daily[d][1] for d in chart_labels]

        daily_docs = defaultdict(lambda: [0, 0])
        for r in doc_rows:
            daily_docs[str(r.date)][0] += r.plan_docs or 0
            daily_docs[str(r.date)][1] += r.fact_docs or 0
        doc_labels = sorted(daily_docs.keys())
        doc_plan_arr = [daily_docs[d][0] for d in doc_labels]
        doc_fact_arr = [daily_docs[d][1] for d in doc_labels]

        ctx = {
            "request": request,
            "start": start, "end": end,
            "prep_total": prep_total, "prep_plan": prep_plan, "prep_fact": prep_fact,
            "prep_fact_count": len(prep_fact_rows), "prep_completion": prep_completion,
            "prep_partial": prep_partial, "prep_issues": prep_issues,
            "weld_total": weld_total, "weld_need": weld_need, "weld_issues": weld_issues,
            "doc_total": doc_total, "doc_plan": doc_plan, "doc_fact": doc_fact,
            "doc_completion": doc_completion, "doc_issues": doc_issues,
            "chart_labels": json.dumps(chart_labels),
            "chart_plan": json.dumps(chart_plan),
            "chart_fact": json.dumps(chart_fact),
            "doc_labels": json.dumps(doc_labels),
            "doc_plan_arr": json.dumps(doc_plan_arr),
            "doc_fact_arr": json.dumps(doc_fact_arr),
        }
        return templates.TemplateResponse("dashboard.html", ctx)
    finally:
        s.close()

import csv, io
from fastapi.responses import StreamingResponse

@app.get("/export")
def export_csv(request: Request, start: str = "", end: str = ""):
    if not is_ceo(request):
        raise HTTPException(403, "Нет доступа")
    # reuse dashboard logic by recalculating
    s = db()
    try:
        def parse(x):
            try: return dt_date.fromisoformat(x)
            except Exception: return None
        d0, d1 = parse(start), parse(end)
        def filt(q, model):
            if d0: q = q.filter(model.date >= d0)
            if d1: q = q.filter(model.date <= d1)
            return q
        prep_rows = filt(s.query(PrepRun), PrepRun).all()
        weld_rows = filt(s.query(WeldNeed), WeldNeed).all()
        doc_rows = filt(s.query(DocRun), DocRun).all()

        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(["Участок", "Записей", "План", "Факт", "Выполнение", "Проблем"])
        prep_plan = sum((r.plan_qty or 0) for r in prep_rows)
        prep_fact = sum(r.fact_qty for r in prep_rows if r.fact_qty is not None)
        prep_comp = (prep_fact / prep_plan) if prep_plan else ""
        w.writerow(["Заготовка", len(prep_rows), prep_plan, prep_fact,
                    (f"{prep_comp*100:.1f}%" if prep_comp != "" else "н/д"),
                    s.query(DataIssue).filter_by(table_name="prep_runs", resolved=False).count()])
        w.writerow(["Потребность сварки", len(weld_rows), "", sum((r.need_qty or 0) for r in weld_rows), "",
                    s.query(DataIssue).filter_by(table_name="weld_needs", resolved=False).count()])
        doc_plan = sum((r.plan_docs or 0) for r in doc_rows)
        doc_fact = sum((r.fact_docs or 0) for r in doc_rows)
        doc_comp = (doc_fact / doc_plan) if doc_plan else ""
        w.writerow(["Документация", len(doc_rows), doc_plan, doc_fact,
                    (f"{doc_comp*100:.1f}%" if doc_comp != "" else "н/д"),
                    s.query(DataIssue).filter_by(table_name="docs_runs", resolved=False).count()])
        # detail prep
        w.writerow([])
        w.writerow(["Детализация: Заготовка"])
        w.writerow(["record_id","date","batch_id","employee","plan_qty","fact_qty"])
        for r in prep_rows:
            w.writerow([r.record_id, r.date, r.batch_id, r.employee, r.plan_qty, r.fact_qty])
        w.writerow([])
        w.writerow(["Детализация: Потребность сварки"])
        w.writerow(["record_id","date","batch_id","need_qty"])
        for r in weld_rows:
            w.writerow([r.record_id, r.date, r.batch_id, r.need_qty])
        w.writerow([])
        w.writerow(["Детализация: Документация"])
        w.writerow(["record_id","date","employee","plan_docs","fact_docs"])
        for r in doc_rows:
            w.writerow([r.record_id, r.date, r.employee, r.plan_docs, r.fact_docs])
        # issues
        w.writerow([])
        w.writerow(["Проблемы данных"])
        w.writerow(["table","record_id","field","issue_type","severity","message","resolved"])
        for i in s.query(DataIssue).all():
            w.writerow([i.table_name, i.record_id, i.field, i.issue_type, i.severity, i.message, i.resolved])
        buf.seek(0)
        return StreamingResponse(iter([buf.getvalue()]), media_type="text/csv",
                                 headers={"Content-Disposition": "attachment; filename=pbs_report.csv"})
    finally:
        s.close()


from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

@app.get("/export.xlsx")
def export_xlsx(request: Request, start: str = "", end: str = ""):
    if not is_ceo(request):
        raise HTTPException(403, "Нет доступа")
    s = db()
    try:
        def parse(x):
            try: return dt_date.fromisoformat(x)
            except Exception: return None
        d0, d1 = parse(start), parse(end)
        def filt(q, model):
            if d0: q = q.filter(model.date >= d0)
            if d1: q = q.filter(model.date <= d1)
            return q
        prep_rows = filt(s.query(PrepRun), PrepRun).all()
        weld_rows = filt(s.query(WeldNeed), WeldNeed).all()
        doc_rows = filt(s.query(DocRun), DocRun).all()
        issues = s.query(DataIssue).all()

        wb = Workbook()
        thin = Side(style="thin", color="BBBBBB")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        blue = PatternFill("solid", fgColor="1F4E79")
        teal = PatternFill("solid", fgColor="1F7A8C")
        purple = PatternFill("solid", fgColor="5B3A8C")
        orange = PatternFill("solid", fgColor="C55A11")
        hdr_font = Font(color="FFFFFF", bold=True)

        def style_header(ws, row, fill, ncols):
            for c in range(1, ncols+1):
                cell = ws.cell(row=row, column=c)
                cell.fill = fill
                cell.font = hdr_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border
            ws.row_dimensions[row].height = 30

        def widths(ws, widths_list):
            for i, w in enumerate(widths_list, start=1):
                ws.column_dimensions[get_column_letter(i)].width = min(w, 45)

        # ---- Сводка ----
        ws = wb.active
        ws.title = "Сводка"
        ws["A1"] = "Отчёт PBS — панель генерального директора"
        ws["A1"].font = Font(size=14, bold=True)
        ws["A2"] = f"Период: {start or 'все данные'} — {end or 'все данные'}"
        ws["A2"].font = Font(italic=True)

        # KPI blocks
        r = 4
        ws.cell(row=r, column=1, value="Участок").font = Font(bold=True)
        ws.cell(row=r, column=2, value="Записей").font = Font(bold=True)
        ws.cell(row=r, column=3, value="План").font = Font(bold=True)
        ws.cell(row=r, column=4, value="Факт").font = Font(bold=True)
        ws.cell(row=r, column=5, value="Выполнение").font = Font(bold=True)
        ws.cell(row=r, column=6, value="Проблем").font = Font(bold=True)
        style_header(ws, r, blue, 6)

        prep_plan = sum((x.plan_qty or 0) for x in prep_rows)
        prep_fact_rows = [x for x in prep_rows if x.fact_qty is not None]
        prep_fact = sum(x.fact_qty for x in prep_fact_rows)
        prep_comp = (prep_fact / prep_plan) if prep_plan else None
        prep_issues = s.query(DataIssue).filter_by(table_name="prep_runs", resolved=False).count()

        weld_need = sum((x.need_qty or 0) for x in weld_rows)
        weld_issues = s.query(DataIssue).filter_by(table_name="weld_needs", resolved=False).count()

        doc_plan = sum((x.plan_docs or 0) for x in doc_rows)
        doc_fact = sum((x.fact_docs or 0) for x in doc_rows)
        doc_comp = (doc_fact / doc_plan) if doc_plan else None
        doc_issues = s.query(DataIssue).filter_by(table_name="docs_runs", resolved=False).count()

        rows_data = [
            ("Заготовка", len(prep_rows), prep_plan, prep_fact,
             (prep_comp if prep_comp is not None else None), prep_issues,
             ("Частичный" if len(prep_fact_rows) < len(prep_rows) else "Полный")),
            ("Потребность сварки", len(weld_rows), None, weld_need, None, weld_issues, "—"),
            ("Документация", len(doc_rows), doc_plan, doc_fact,
             (doc_comp if doc_comp is not None else None), doc_issues, "Полный"),
        ]
        rr = r + 1
        for name, cnt, plan, fact, comp, iss, status in rows_data:
            ws.cell(row=rr, column=1, value=name)
            ws.cell(row=rr, column=2, value=cnt)
            if plan is not None: ws.cell(row=rr, column=3, value=plan)
            if fact is not None:
                ws.cell(row=rr, column=4, value=fact)
            if comp is not None:
                c = ws.cell(row=rr, column=5, value=comp); c.number_format = "0.0%"
            ws.cell(row=rr, column=6, value=iss)
            ws.cell(row=rr, column=7, value=status)
            for cc in range(1, 8):
                ws.cell(row=rr, column=cc).border = border
            rr += 1
        ws.cell(row=r, column=7, value="Статус").font = Font(bold=True)
        style_header(ws, r, blue, 7)
        widths(ws, [24, 12, 12, 14, 14, 10, 14])

        # reliability notes
        rr += 1
        if len(prep_fact_rows) < len(prep_rows):
            ws.cell(row=rr, column=1, value=f"Заготовка: факт рассчитан по {len(prep_fact_rows)} из {len(prep_rows)} записей. Отсутствующие значения не заменены нулями.")
            ws.cell(row=rr, column=1).font = Font(color="C55A11", italic=True)
            rr += 1
        if prep_issues or weld_issues or doc_issues:
            ws.cell(row=rr, column=1, value="Обнаружены подозрительные значения. Показатели помечены как частичные/требующие проверки.")
            ws.cell(row=rr, column=1).font = Font(color="C55A11", italic=True)

        # ---- detail sheets ----
        def write_table(sheetname, rows, headers, fill, widths_list, fmts=None):
            wsh = wb.create_sheet(sheetname)
            wsh.append(headers)
            style_header(wsh, 1, fill, len(headers))
            for row in rows:
                wsh.append(row)
            ws_last = wsh.max_row
            for rr in range(2, ws_last+1):
                for cc in range(1, len(headers)+1):
                    c = wsh.cell(row=rr, column=cc)
                    c.border = border
                    if fmts and fmts.get(cc):
                        c.number_format = fmts[cc]
                if rr % 2 == 0:
                    for cc in range(1, len(headers)+1):
                        wsh.cell(row=rr, column=cc).fill = PatternFill("solid", fgColor="F2F2F2")
            wsh.freeze_panes = "A2"
            wsh.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws_last}"
            widths(wsh, widths_list)
            # print settings
            wsh.page_setup.orientation = "landscape"
            wsh.sheet_properties.pageSetUpPr.fitToPage = True
            wsh.page_setup.fitToWidth = 1
            wsh.print_title_rows = "1:1"
            return wsh

        fmt_date = "yyyy-mm-dd"
        prep_data = [[r.record_id, r.date, r.batch_id, r.employee, r.plan_qty, r.fact_qty] for r in prep_rows]
        write_table("Заготовка", prep_data,
                    ["record_id", "date", "batch_id", "employee", "plan_qty", "fact_qty"],
                    blue, [14, 12, 12, 20, 10, 10],
                    fmts={2: fmt_date})

        weld_data = [[r.record_id, r.date, r.batch_id, r.need_qty] for r in weld_rows]
        write_table("Потребность сварки", weld_data,
                    ["record_id", "date", "batch_id", "need_qty"],
                    teal, [14, 12, 12, 12], fmts={2: fmt_date})

        doc_data = [[r.record_id, r.date, r.employee, r.plan_docs, r.fact_docs] for r in doc_rows]
        write_table("Документация", doc_data,
                    ["record_id", "date", "employee", "plan_docs", "fact_docs"],
                    purple, [14, 12, 20, 10, 10], fmts={2: fmt_date})

        # problems sheet with severity colors
        ws_p = wb.create_sheet("Проблемы данных")
        headers = ["Участок", "ID записи", "Поле", "Тип проблемы", "Серьёзность", "Описание", "Исправлено"]
        ws_p.append(headers)
        style_header(ws_p, 1, orange, len(headers))
        sev_fill = {"high": PatternFill("solid", fgColor="FFC7CE"),
                    "medium": PatternFill("solid", fgColor="FFEB9C"),
                    "low": PatternFill("solid", fgColor="C6EFCE")}
        sev_ru = {"high": "Высокая", "medium": "Средняя", "low": "Низкая"}
        type_ru = {
            "duplicate_record": "Дубликат записи",
            "missing_value": "Отсутствующее значение",
            "invalid_value": "Некорректное значение",
            "negative_value": "Отрицательное значение",
            "implausible_value": "Подозрительное значение",
            "future_date": "Дата в будущем",
        }
        table_ru = {"prep_runs": "Заготовка", "weld_needs": "Потребность сварки", "docs_runs": "Документация"}
        for i in issues:
            ws_p.append([table_ru.get(i.table_name, i.table_name),
                         i.record_id, i.field, type_ru.get(i.issue_type, i.issue_type),
                         sev_ru.get(i.severity, i.severity), i.message, "Да" if i.resolved else "Нет"])
        for rr in range(2, ws_p.max_row + 1):
            sev = ws_p.cell(row=rr, column=5).value
            rev = {"Высокая": "high", "Средняя": "medium", "Низкая": "low"}
            fill = sev_fill.get(rev.get(sev))
            for cc in range(1, len(headers) + 1):
                c = ws_p.cell(row=rr, column=cc)
                c.border = border
                if fill:
                    c.fill = fill
        ws_p.freeze_panes = "A2"
        ws_p.auto_filter.ref = f"A1:G{ws_p.max_row}"
        widths(ws_p, [16, 12, 14, 20, 10, 45, 10])
        ws_p.page_setup.orientation = "landscape"
        ws_p.sheet_properties.pageSetUpPr.fitToPage = True
        ws_p.page_setup.fitToWidth = 1
        ws_p.print_title_rows = "1:1"

        # summary print settings
        ws.page_setup.orientation = "landscape"

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": "attachment; filename=pbs_report.xlsx"})
    finally:
        s.close()
app.add_middleware(SessionMiddleware, secret_key="pbs-demo-secret")
